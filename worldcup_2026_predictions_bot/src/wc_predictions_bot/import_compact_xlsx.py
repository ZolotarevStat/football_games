from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from .config import Config
from .env_loader import load_dotenv
from .setup_sheets import build_service
from .sheet_schema import SHEET_HEADERS


TZ = ZoneInfo("Europe/Moscow")
TEAM_CODES = {
    "Австралия": "Aus",
    "Австрия": "Aut",
    "Алжир": "Alg",
    "Англия": "Eng",
    "Аргентина": "Arg",
    "Бельгия": "Bel",
    "Босния и Герцеговина": "BiH",
    "Бразилия": "Bra",
    "Гаити": "Hai",
    "Гана": "Gha",
    "Германия": "Ger",
    "ДР Конго": "DRC",
    "Египет": "Egy",
    "Иордания": "Jor",
    "Ирак": "Irq",
    "Иран": "Irn",
    "Испания": "Esp",
    "Кабо-Верде": "CPV",
    "Канада": "Can",
    "Катар": "Qat",
    "Колумбия": "Col",
    "Кот-д'Ивуар": "CIV",
    "Кюрасао": "Cur",
    "Марокко": "Mor",
    "Мексика": "Mex",
    "Нидерланды": "Ned",
    "Новая Зеландия": "NZl",
    "Норвегия": "Nor",
    "Панама": "Pan",
    "Парагвай": "Par",
    "Португалия": "Por",
    "США": "USA",
    "Саудовская Аравия": "KSA",
    "Сенегал": "Sen",
    "Тунис": "Tun",
    "Турция": "Tur",
    "Узбекистан": "Uzb",
    "Уругвай": "Uru",
    "Франция": "Fra",
    "Хорватия": "Cro",
    "Чехия": "Cze",
    "Швейцария": "Sui",
    "Швеция": "Swe",
    "Шотландия": "Sco",
    "Эквадор": "Ecu",
    "ЮАР": "SAf",
    "Южная Корея": "SKo",
    "Япония": "Jpn",
}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--env-file", default=".env")
    parser.add_argument("--spreadsheet-id", default="")
    parser.add_argument("--service-account-json-b64", default="")
    parser.add_argument("--service-account-file", default="")
    args = parser.parse_args()

    load_dotenv(args.env_file)
    config = Config.from_env()
    spreadsheet_id = args.spreadsheet_id or config.spreadsheet_id
    service_account_json_b64 = args.service_account_json_b64 or config.service_account_json_b64
    service_account_file = args.service_account_file or config.service_account_file
    if not spreadsheet_id:
        raise RuntimeError("GOOGLE_SPREADSHEET_ID is empty.")

    matches, players = parse_workbook(Path(args.workbook))
    service = build_service(service_account_json_b64, service_account_file)
    values = service.spreadsheets().values()
    replace_sheet(values, spreadsheet_id, "matches", matches)
    replace_sheet(values, spreadsheet_id, "players", players)
    print(f"Imported matches={len(matches)} players={len(players)} into spreadsheet.")


def parse_workbook(path: Path) -> tuple[list[list[str]], list[list[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required only for this import script.") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    matches = parse_matches(workbook["Матчи"])
    players = parse_players(workbook["Игроки"])
    return matches, players


def parse_matches(sheet: Any) -> list[list[str]]:
    rows: list[list[str]] = []
    for index, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), start=1):
        group, tour, _, kickoff, deadline, match_name = row[:6]
        if not match_name:
            continue
        team1, team2 = parse_match_name(str(match_name))
        rows.append(
            [
                match_id(team1, team2),
                str(group or "").strip(),
                str(tour or "").strip(),
                iso_msk(kickoff),
                iso_msk(deadline),
                team1,
                team2,
                "open",
            ]
        )
    return rows


def parse_players(sheet: Any) -> list[list[str]]:
    raw_players: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=5, values_only=True):
        team, position, player_name, club, league, _, weighted_ga, national_ga_2y, _, _, _, priority = row[:12]
        if not team or not player_name:
            continue
        raw_players.append(
            {
                "team": normalize_team(str(team)),
                "position": str(position or "").strip(),
                "player_name_ru": str(player_name or "").strip(),
                "player_name_en": "",
                "club": str(club or "").strip(),
                "league": str(league or "").strip(),
                "weighted_ga": format_number(weighted_ga),
                "national_ga_2y": format_number(national_ga_2y),
                "priority": numeric(priority),
            }
        )

    by_team: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for player in raw_players:
        by_team[player["team"]].append(player)

    rows: list[list[str]] = []
    for team in sorted(by_team):
        team_players = sorted(
            by_team[team],
            key=lambda player: (-player["priority"], player["position"], player["player_name_ru"]),
        )
        for rank, player in enumerate(team_players, start=1):
            rows.append(
                [
                    player["team"],
                    player["position"],
                    player["player_name_ru"],
                    player["player_name_en"],
                    player["club"],
                    player["league"],
                    player["weighted_ga"],
                    player["national_ga_2y"],
                    format_number(player["priority"]),
                    str(rank),
                    "TRUE",
                ]
            )
    return rows


def replace_sheet(values: Any, spreadsheet_id: str, sheet_name: str, rows: list[list[str]]) -> None:
    headers = SHEET_HEADERS[sheet_name]
    values.clear(spreadsheetId=spreadsheet_id, range=f"{sheet_name}!A:Z").execute()
    values.update(
        spreadsheetId=spreadsheet_id,
        range=f"{sheet_name}!A1",
        valueInputOption="USER_ENTERED",
        body={"values": [headers, *rows]},
    ).execute()


def parse_match_name(match_name: str) -> tuple[str, str]:
    if " - " not in match_name:
        raise ValueError(f"Match name must contain ' - ': {match_name}")
    team1, team2 = match_name.split(" - ", 1)
    return normalize_team(team1), normalize_team(team2)


def match_id(team1: str, team2: str) -> str:
    try:
        return f"{TEAM_CODES[team1]}{TEAM_CODES[team2]}"
    except KeyError as exc:
        raise ValueError(f"Missing team code for {exc.args[0]}") from exc


def normalize_team(value: str) -> str:
    return value.strip().replace("’", "'").replace("`", "'")


def iso_msk(value: Any) -> str:
    if not isinstance(value, datetime):
        raise ValueError(f"Expected datetime, got {value!r}")
    if value.tzinfo is None:
        value = value.replace(tzinfo=TZ)
    return value.isoformat(timespec="seconds")


def numeric(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    return float(str(value).replace(",", "."))


def format_number(value: Any) -> str:
    number = numeric(value)
    if number.is_integer():
        return str(int(number))
    return f"{number:.3f}".rstrip("0").rstrip(".")


if __name__ == "__main__":
    main()
